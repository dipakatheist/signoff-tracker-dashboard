from __future__ import annotations

import argparse
import json
import threading
import webbrowser
from collections import Counter, defaultdict
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOKS = {
    "pkg1-2": Path.home() / "Desktop" / "Sept'26 PKG1-2 Fresh-Resignof Signoff Sheet.xlsx",
    "pkg-7": Path.home() / "Desktop" / "PKG-7 (September) Fresh-Resignoff Pending Signof Sheet as on 08-09-26.xlsx",
}
PACKAGE_LABELS = {"pkg1-2": "PKG1-2", "pkg-7": "PKG-7"}
HEADER_ALIASES = {
    "division": ("sub division", "subdivision"),
    "executive": ("excutive", "exceutive", "executive"),
    "kind": ("fresh/re-signoff", "month of signoff"),
    "mail": ("mail status",),
    "signoff": ("signoff-status", "signoff status"),
    "communication": ("comm status",),
    "priority": ("priority",),
    "date": ("signoff date",),
}

_cache_lock = threading.Lock()
_cached_signature: tuple[tuple[str, int, int], ...] | None = None
_cached_snapshot: dict | None = None


def normalize(value: object) -> str:
    return " ".join(str(value or "").split()).casefold()


def column_indexes(header: tuple) -> dict[str, int]:
    available = {}
    for index, value in enumerate(header):
        available.setdefault(normalize(value), index)
    result = {}
    for name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in available:
                result[name] = available[alias]
                break
    missing = set(HEADER_ALIASES) - {"communication"} - set(result)
    if missing:
        raise ValueError("The Data sheet is missing required dashboard columns.")
    return result


def value_at(row: tuple, indexes: dict[str, int], key: str) -> object:
    index = indexes.get(key)
    if index is None:
        return None
    return row[index] if index < len(row) else None


def record_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def new_bucket() -> dict:
    return {
        "total": 0,
        "done": 0,
        "pending": 0,
        "mailPending": 0,
        "fresh": 0,
        "resignoff": 0,
        "otherType": 0,
        "statusCounts": Counter(),
        "mailCounts": Counter(),
        "communicationCounts": Counter(),
        "priorityCounts": Counter(),
        "daily": Counter(),
        "divisions": defaultdict(Counter),
        "executives": defaultdict(Counter),
    }


def update_group(bucket: dict, row: tuple, indexes: dict[str, int], kind: str | None) -> None:
    bucket["total"] += 1
    bucket[kind if kind in ("fresh", "resignoff") else "otherType"] += 1

    signoff = str(value_at(row, indexes, "signoff") or "Unspecified").strip()
    mail = str(value_at(row, indexes, "mail") or "Unspecified").strip()
    communication = str(value_at(row, indexes, "communication") or "Unspecified").strip()
    priority = str(value_at(row, indexes, "priority") or "Unspecified").strip()
    bucket["statusCounts"][signoff] += 1
    bucket["mailCounts"][mail] += 1
    bucket["communicationCounts"][communication] += 1
    bucket["priorityCounts"][priority] += 1

    status = normalize(signoff)
    if status == "done":
        bucket["done"] += 1
    elif "pending" in status:
        bucket["pending"] += 1
    if "pending" in normalize(mail):
        bucket["mailPending"] += 1

    subdivision = str(value_at(row, indexes, "division") or "Unassigned").strip()
    executive = str(value_at(row, indexes, "executive") or "Unassigned").strip()
    for key, label in (("divisions", subdivision), ("executives", executive)):
        item = bucket[key][label]
        item["total"] += 1
        if status == "done":
            item["done"] += 1
        elif "pending" in status:
            item["pending"] += 1

    signed_on = record_date(value_at(row, indexes, "date"))
    if signed_on:
        bucket["daily"][signed_on.isoformat()] += 1


def breakdown(source: defaultdict) -> list[dict]:
    items = []
    for name, values in source.items():
        classified = values["done"] + values["pending"]
        items.append({
            "name": name,
            "total": values["total"],
            "done": values["done"],
            "pending": values["pending"],
            "completion": round(values["done"] / classified * 100, 1) if classified else 0,
        })
    return sorted(items, key=lambda item: (-item["pending"], -item["total"], item["name"].casefold()))


def finish_bucket(bucket: dict) -> dict:
    classified = bucket["done"] + bucket["pending"]
    return {
        "summary": {
            "total": bucket["total"],
            "done": bucket["done"],
            "pending": bucket["pending"],
            "mailPending": bucket["mailPending"],
            "unclassified": bucket["total"] - classified,
            "completion": round(bucket["done"] / classified * 100, 1) if classified else 0,
        },
        "typeCounts": {
            "fresh": bucket["fresh"],
            "resignoff": bucket["resignoff"],
            "other": bucket["otherType"],
        },
        "signoffStatuses": bucket["statusCounts"].most_common(),
        "mailStatuses": bucket["mailCounts"].most_common(),
        "communicationStatuses": bucket["communicationCounts"].most_common(),
        "priorities": bucket["priorityCounts"].most_common(),
        "dailySignoffs": sorted(bucket["daily"].items()),
        "divisions": breakdown(bucket["divisions"]),
        "executives": breakdown(bucket["executives"]),
    }


def merge_breakdowns(groups: list[dict], key: str) -> list[dict]:
    merged = defaultdict(lambda: {"total": 0, "done": 0, "pending": 0})
    for group in groups:
        for item in group[key]:
            target = merged[item["name"]]
            for field in ("total", "done", "pending"):
                target[field] += item[field]
    return breakdown(merged)


def combine_groups(groups: list[dict]) -> dict:
    summary = {
        key: sum(group["summary"][key] for group in groups)
        for key in ("total", "done", "pending", "mailPending", "unclassified")
    }
    classified = summary["done"] + summary["pending"]
    summary["completion"] = round(summary["done"] / classified * 100, 1) if classified else 0
    type_counts = {
        key: sum(group["typeCounts"][key] for group in groups)
        for key in ("fresh", "resignoff", "other")
    }

    def merge_counts(key: str) -> list[tuple[str, int]]:
        counts = Counter()
        for group in groups:
            counts.update(dict(group[key]))
        return counts.most_common()

    daily = Counter()
    for group in groups:
        daily.update(dict(group["dailySignoffs"]))
    return {
        "summary": summary,
        "typeCounts": type_counts,
        "signoffStatuses": merge_counts("signoffStatuses"),
        "mailStatuses": merge_counts("mailStatuses"),
        "communicationStatuses": merge_counts("communicationStatuses"),
        "priorities": merge_counts("priorities"),
        "dailySignoffs": sorted(daily.items()),
        "divisions": merge_breakdowns(groups, "divisions"),
        "executives": merge_breakdowns(groups, "executives"),
    }


def read_dashboard(workbook_path: Path, modified_at: datetime) -> dict:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Data" not in workbook.sheetnames:
            raise ValueError("The workbook does not contain a Data sheet.")
        sheet = workbook["Data"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("The Data sheet is empty.")
        indexes = column_indexes(header)
        buckets = {"all": new_bucket(), "fresh": new_bucket(), "resignoff": new_bucket()}
        latest_signoff = None

        for row in rows:
            if not any(value not in (None, "") for value in row):
                continue
            category = normalize(value_at(row, indexes, "kind"))
            kind = "fresh" if "fresh" in category else "resignoff" if "re-signoff" in category or "resignoff" in category else None
            targets = [buckets["all"]]
            if kind:
                targets.append(buckets[kind])
            for target in targets:
                update_group(target, row, indexes, kind)
            signed_on = record_date(value_at(row, indexes, "date"))
            if signed_on and (latest_signoff is None or signed_on > latest_signoff):
                latest_signoff = signed_on

        return {
            "source": workbook_path.name,
            "sourceModifiedAt": modified_at.astimezone().isoformat(timespec="seconds"),
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "latestSignoffDate": latest_signoff.isoformat() if latest_signoff else None,
            "groups": {key: finish_bucket(value) for key, value in buckets.items()},
        }
    finally:
        workbook.close()


def get_snapshot(workbook_paths: dict[str, Path]) -> dict:
    global _cached_signature, _cached_snapshot
    with _cache_lock:
        try:
            signatures = []
            stats = {}
            for package, workbook_path in sorted(workbook_paths.items()):
                stat = workbook_path.stat()
                signatures.append((package, stat.st_mtime_ns, stat.st_size))
                stats[package] = stat
            signature = tuple(signatures)
            if _cached_snapshot is not None and signature == _cached_signature:
                return _cached_snapshot

            package_data = {}
            for package, workbook_path in sorted(workbook_paths.items()):
                stat = stats[package]
                dataset = read_dashboard(workbook_path, datetime.fromtimestamp(stat.st_mtime).astimezone())
                dataset["label"] = PACKAGE_LABELS.get(package, package.upper())
                package_data[package] = dataset
            datasets = list(package_data.values())
            combined = {
                "sourceModifiedAt": max(item["sourceModifiedAt"] for item in datasets),
                "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
                "latestSignoffDate": max((item["latestSignoffDate"] for item in datasets if item["latestSignoffDate"]), default=None),
                "groups": {
                    group: combine_groups([item["groups"][group] for item in datasets])
                    for group in ("all", "fresh", "resignoff")
                },
            }
            snapshot = {"packages": package_data, "combined": combined, "generatedAt": combined["generatedAt"]}
            _cached_signature = signature
            _cached_snapshot = snapshot
            return snapshot
        except Exception as error:
            if _cached_snapshot is not None:
                return {**_cached_snapshot, "warning": "The latest workbook change could not be read. Showing the last successfully loaded data."}
            raise RuntimeError("Could not read the workbook. Check that it exists, is synced locally, and is not being replaced during the refresh.") from error


class DashboardHandler(BaseHTTPRequestHandler):
    def send_body(self, status: int, body: bytes, content_type: str) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        path = urlsplit(self.path).path
        if path in ("/", "/index.html"):
            try:
                body = (ROOT / "index.html").read_bytes()
            except OSError:
                self.send_body(500, b"Dashboard page is unavailable.", "text/plain; charset=utf-8")
                return
            self.send_body(200, body, "text/html; charset=utf-8")
            return
        if path == "/api/dashboard":
            try:
                body = json.dumps(get_snapshot(self.server.workbook_paths), ensure_ascii=False).encode("utf-8")
                self.send_body(200, body, "application/json; charset=utf-8")
            except RuntimeError as error:
                body = json.dumps({"error": str(error)}).encode("utf-8")
                self.send_body(503, body, "application/json; charset=utf-8")
            return
        if path == "/health":
            self.send_body(200, b"ok", "text/plain; charset=utf-8")
            return
        self.send_body(404, b"Not found.", "text/plain; charset=utf-8")

    def log_message(self, format: str, *args: object) -> None:
        return


def main() -> None:
    parser = argparse.ArgumentParser(description="Serve a local, aggregate-only signoff dashboard for multiple packages.")
    parser.add_argument("--pkg1-2", dest="pkg1_2", type=Path, default=DEFAULT_WORKBOOKS["pkg1-2"])
    parser.add_argument("--pkg-7", dest="pkg_7", type=Path, default=DEFAULT_WORKBOOKS["pkg-7"])
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()

    server = ThreadingHTTPServer((args.host, args.port), DashboardHandler)
    server.workbook_paths = {
        "pkg1-2": args.pkg1_2.expanduser(),
        "pkg-7": args.pkg_7.expanduser(),
    }
    address = f"http://{args.host}:{args.port}"
    print(f"Renewal dashboard: {address}")
    for package, workbook_path in server.workbook_paths.items():
        print(f"{PACKAGE_LABELS[package]} workbook: {workbook_path.name}")
    print("Only aggregate dashboard metrics are served. Press Ctrl+C to stop.")
    threading.Timer(0.8, webbrowser.open, args=(address,)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nDashboard stopped.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()